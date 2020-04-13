#!/usr/bin/env python3


#
# Copyright (C) 2018 ETH Zurich and University of Bologna
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
#


from prettytable import PrettyTable
import os
import re
from xml.sax.saxutils import escape
#import plpdb
import datetime
#import plptree
import subprocess

try:
    import smtplib
    from openpyxl import Workbook
    from xlsxwriter.utility import xl_rowcol_to_cell
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.base import MIMEBase
    from email import encoders

except:
    raise


class Package_Build(object):

    def __init__(self, pobj, package, config=None, status=None, start_date=None,
                 end_date=None, version=None, artifact=None, deps=None,
                passed=None, tests=None, testrun=None, id=None, branch=None,
                commit_version=None, env=[]):
        self.id = id
        self.status = status
        self.start_date = start_date
        self.end_date = end_date
        self.package = package
        self.pobj = pobj
        self.version = version
        self.artifact = artifact
        self.deps = deps
        self.tests = tests
        self.passed = passed
        self.testrun = testrun
        self.config = config
        self.dbobj = None
        self.branch = branch
        self.commit_version = commit_version
        self.env = env

    def set_status(self, status):
        self.status = status

    def set_start_date(self, start_date):
        self.start_date = start_date

    def set_end_date(self, end_date):
        self.end_date = end_date

    def set_nb_tests(self, tests):
        self.tests = tests

    def set_nb_passed(self, tests):
        self.passed = tests

    def commit(self):
        if self.pobj.db is not None:
            if self.dbobj is None:
                self.dbobj = self.pobj.db.reg_build(
                    package=self.package, status=self.status,
                    start_date=self.start_date, end_date=self.end_date,
                    version=self.version, artifact=self.artifact,
                    deps=" ".join(self.deps), passed=self.passed,
                    tests=self.tests, testrun=self.testrun, config=self.config,
                    branch=self.branch, commit=self.commit_version,
                    env=self.env)

                if self.pobj.db_info is not None:
                    with open(self.pobj.db_info, "a") as file:
                        file.write('%s.build.id=%d\n' % (self.package, self.dbobj.id))

            else:
                self.dbobj.package = self.package
                self.dbobj.status = self.status
                self.dbobj.start_date = self.start_date
                self.dbobj.end_date = self.end_date
                self.dbobj.version = self.version
                self.dbobj.artifact = self.artifact
                self.dbobj.deps = " ".join(self.deps)
                self.dbobj.passed = self.passed
                self.dbobj.tests = self.tests
                self.dbobj.testrun = self.testrun
                self.dbobj.config = self.config
                self.dbobj.branch = self.branch
                self.dbobj.commit = self.commit_version
                self.pobj.db.session.commit()

        return self.dbobj


class PulpSdk_build(object):
    def __init__(self, sdk, distrib, status):
        self.status = status
        self.distrib = distrib
        sdk.regBuild(self)

class PulpSdk(object):
    def __init__(self, version=None):
        self.builds = {}
        self.setVersion(version)

    def regBuild(self, build):
        if self.builds.get(build.distrib) is None:
            self.builds[build.distrib] = []
        self.builds[build.distrib].append(build)

    def setVersion(self, version):
        if self.version is not None:
            del pulp_sdks[self.version]
        self.version = version
        pulp_sdks[version] = self


class Test(object):
    def __init__(self, name, parent=None):
        self.parent = parent
        self.name = name
        self.runs = []
        self.tests = []
        if parent is not None:
            parent.regTest(self)
        self.configs = []
        self.runs_for_config = {}


    def get_run_for_build(self, config, build):
        config_runs = self.runs_for_config.get(config)
        if config_runs is not None:
            for run in config_runs:
                if run.build == build:
                    return run

        return None


    def dump_reg(self, config, ref_build, build, table):

        nb_new = 0
        nb_removed = 0
        nb_regs = 0
        nb_fixed = 0

        for test in self.tests:
            (test_regs, test_fixed, test_removed, test_new) = \
                test.dump_reg(config, ref_build, build, table)
            nb_new += test_new
            nb_removed += test_removed
            nb_regs += test_regs
            nb_fixed += test_fixed

        ref_run = self.get_run_for_build(config, ref_build)
        run = self.get_run_for_build(config, build)

        if ref_run is None:
            if run is not None:
                nb_new += 1
        elif run is None:
            if ref_run is not None:
                nb_removed += 1
        elif ref_run.success:
            if not run.success:
                nb_regs += 1
        elif not ref_run.success:
            if run.success:
                nb_fixed += 1

        row = [self.name]

        if nb_regs + nb_fixed + nb_removed + nb_new != 0:
            row.append('%d/%d/%d/%d' % (nb_regs, nb_fixed, nb_removed, nb_new))

            table.append(row)

        return (nb_regs, nb_fixed, nb_removed, nb_new)

    def check_reg(self, config, ref_build, build):
        ref_build_status = None
        build_status = None

        config_runs = self.runs_for_config.get(config)
        if config_runs is not None:
            for run in config_runs:
                if run.build == ref_build:
                    ref_build_status = run.success
                elif run.build == build:
                    build_status = run.success

            if ref_build_status is not None and \
                build_status is not None and \
                ref_build_status != build_status:
                return False

        for test in self.tests:
            if not test.check_reg(config, ref_build, build):
                return False

        return True

    def regTest(self, test):
        self.tests.append(test)

    def get_configs(self):
        return self.configs

    def reg_config(self, config):
        if not config in self.configs:
            self.configs.append(config)
        if self.parent is not None:
            self.parent.reg_config(config)

    def regRun(self, run):
        self.runs.append(run)
        self.reg_config(run.config)
        if self.runs_for_config.get(run.config) is None:
            self.runs_for_config[run.config] = []
        self.runs_for_config[run.config].append(run)

    def getNbTests(self, config=None, build=None):
        nbTests = 0

        if config is None:
            runs = self.runs
        else:
            runs = self.runs_for_config.get(config)
        if runs is not None:
            for run in runs:
                if build is None or run.build == build:
                    nbTests += 1

        for test in self.tests:
            nbTests += test.getNbTests(config, build)
        return nbTests

    def getNbSuccess(self, config=None, build=None):
        nbSuccess = 0

        if config is None:
            runs = self.runs
        else:
            runs = self.runs_for_config.get(config)

        if runs is not None:
            for run in runs:
                if (build is None or run.build == build) and run.success:
                    nbSuccess += 1

        for test in self.tests:
            nbSuccess += test.getNbSuccess(config, build)
        return nbSuccess

    def getDuration(self):
        duration = 0.0
        for run in self.runs:
            duration += run.duration
        for test in self.tests:
            duration += test.getDuration()
        return duration

    def dump_status(self, config, builds, table):

        row = [self.name]

        for test in self.tests:
            test.dump_status(config, builds, table)

        for build in builds:
            nbTests = self.getNbTests(config=config, build=build)
            nbSuccess = self.getNbSuccess(config=config, build=build)
            row.append('%d/%d' % (nbSuccess, nbTests))


        table.append(row)

    def dumpToConsole(self, table):
        nbTests = self.getNbTests()
        nbSuccess = self.getNbSuccess()

        if len(self.tests) == 0:
            if len(self.runs) == 1:
                table.add_row([
                    self.name, self.runs[0].config,
                    self.runs[0].duration, '%d/%d' % (nbSuccess, nbTests)])
            else:
                table.add_row([
                    self.name, '', '%.2f' % self.getDuration(),
                    '%d/%d' % (nbSuccess, nbTests)])
                for run in self.runs:
                    table.add_row([
                        '', run.config, '%.2f' % run.duration,
                        '%d/%d' % (run.success, 1)])
        else:
            table.add_row([
                self.name, '', '%.2f' % self.getDuration(),
                '%d/%d' % (nbSuccess, nbTests)])
            for test in self.tests:
                test.dumpToConsole(table)

    def dumpToJunit(self, reportPath, testFile=None, test_prefix=''):

        name = self.name.split(':')[-1]

        if test_prefix is None:
            testcase_prefix = name
        else:
            testcase_prefix = test_prefix + ':' + name

        if len(self.tests) != 0:

            if testFile is None:

                fileName = '%s/TEST-%s.xml' % (reportPath, name)
                with open(fileName, 'w') as testFile:
                    nbTests = self.getNbTests()
                    nbSuccess = self.getNbSuccess()
                    nbErrors = nbTests - nbSuccess
                    testFile.write('<?xml version="1.0" encoding="UTF-8"?>\n')
                    testFile.write('<testsuite errors="%d" failures="%d" name="%s" tests="%d" time="%f">\n' % (nbErrors, nbErrors, name, nbTests, self.getDuration()))
                    for testCase in self.tests:
                        testCase.dumpToJunit(reportPath, testFile, testcase_prefix)
                    testFile.write('</testsuite>\n')

            else:
                for testCase in self.tests:
                    testCase.dumpToJunit(reportPath, testFile, testcase_prefix)


        # The report can only be dumped if a testset is launched. TestFile will be None if we are launching a test and not a testset
        if testFile is None:
            return

        for run in self.runs:

            testFile.write('  <testcase classname="%s" name="%s" time="%f">\n' % (run.config, test_prefix + ':' + name, run.duration))
            if run.success:
                testFile.write('    <success/>\n')
            else:
                testFile.write('    <failure>\n')
                for line in run.log:
                    RE_XML_ILLEGAL = u'([\u0000-\u0008\u000b-\u000c\u000e-\u001f\ufffe-\uffff])' + \
                                     u'|' + \
                                     u'([%s-%s][^%s-%s])|([^%s-%s][%s-%s])|([%s-%s]$)|(^[%s-%s])' % \
                                     (chr(0xd800),chr(0xdbff),chr(0xdc00),chr(0xdfff),
                                      chr(0xd800),chr(0xdbff),chr(0xdc00),chr(0xdfff),
                                      chr(0xd800),chr(0xdbff),chr(0xdc00),chr(0xdfff))
                    xmlLine = re.sub(RE_XML_ILLEGAL, "", escape(line))
                    testFile.write(xmlLine)
                testFile.write('</failure>\n')
            testFile.write('  </testcase>\n')

class Config_reg_report(object):

    def __init__(self, nb_fixed, nb_regs, nb_new, nb_removed):
        self.nb_fixed = nb_fixed
        self.nb_regs = nb_regs
        self.nb_new = nb_new
        self.nb_removed = nb_removed


class Reg_report(object):

    def _init__(self):
        self.configs = {}

    def add_config(self, name, reg):
        self.configs[config] = reg

    def dump(self):
        pass

    def has_reg(self):
        return False


class TestRun(object):
    def __init__(self, pobj, test, success, duration, config, log, name=None,
                 metrics='', build=None):
        commit = True

        if test is None:
            # This is the case where the test hierarchy is reconstructed
            # from the database. In this case we need to recreate the test
            # objects
            test = pobj.getTest(name)
            commit = False

        self.test = test
        self.success = success
        self.duration = duration
        self.config = config
        self.log = log
        self.pobj = pobj
        self.status = 'success' if self.success else 'failure'
        self.metrics = metrics
        self.build = build
        if test is not None:
            test.regRun(self)
            if commit:
                self.commit()
        else:
            self.test = Test(name)

    def commit(self):
        if self.pobj.db is not None:
            self.pobj.db.reg_run(
                name=self.test.name,
                build=self.build,
                status=self.status,
                time=datetime.datetime.fromtimestamp(self.duration),
                metrics=self.metrics,
                config=self.config.__str__()
            )


class Report_line(object):

    def __init__(self, line):
        self.line = line

    def get_text(self):
        return self.line + '\n'

    def dump_to_xls(self, ws):
        ws.append([self.line])



class Report_table(object):

    def __init__(self, name, table):
        self.table = table
        self.name = name

    def get_text(self):

        x = PrettyTable(self.table[0])
        x.align[self.table[0][0]] = "l"

        for row in self.table[1:]:
            x.add_row(row)

        return (str(x) + '\n')


    def dump_to_xls(self, ws):

        first_row = ws._current_row

        for row in self.table:
            ws.append(row)

        last_row = ws._current_row - 1

        c0 = xl_rowcol_to_cell(first_row, 0)
        c1 = xl_rowcol_to_cell(last_row, len(self.table[0]) - 1)
        coord = "%s:%s" % (c0, c1)

        tab = Table(displayName=self.name, ref=coord)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
               showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        tab.tableStyleInfo = style
        ws.add_table(tab)


class Report_sheet(object):

    def __init__(self, name, short_name):
        self.objects = []
        self.name = name
        self.short_name = short_name

    def add_line(self, line):
        self.objects.append(Report_line(line))

    def add_table(self, name, table):
        self.objects.append(Report_table(name, table))

    def get_text(self):
        text = self.name + '\n'
        for obj in self.objects:
            text += obj.get_text()
        return text

    def dump_to_xls(self, ws):
        for obj in self.objects:
            obj.dump_to_xls(ws)
        ws.column_dimensions['A'].width = 50



class Report(object):

    def __init__(self):
        self.sheets = []

    def alloc_sheet(self, name, short_name):
        sheet = Report_sheet(name, short_name)
        self.sheets.append(sheet)
        return sheet

    def get_text(self):
        text = ''
        for sheet in self.sheets:
            text += sheet.get_text()
        return text

    def dump_to_console(self):
        print (self.get_text())

    def dump_to_mail(self, subject, attachments=[], author_email=None):

        if author_email is None:
            print ('No email was specified, dropping report')
            return

        print ('Sending report email to: ' + author_email)

        server = smtplib.SMTP('iis-mail.ee.ethz.ch')
         
        msg = MIMEMultipart()
        msg['From'] = "jenkins@iis.ee.ethz.ch"
        msg['To'] = author_email
        msg['Subject'] = subject
         
        body = self.get_text()
        msg.attach(MIMEText(body, 'plain'))

        for attachment in attachments:
            with open(attachment, "rb") as file:
                part = MIMEBase('application', 'octet-stream')
                part.set_payload(file.read())
                encoders.encode_base64(part)
                part.add_header('Content-Disposition', "attachment; filename= %s" % attachment)
                msg.attach(part)
         
        text = msg.as_string()

        server.sendmail("jenkins@iis.ee.ethz.ch", author_email, text)
        server.quit()


    def dump_to_xls(self, xls):

        wb = Workbook()
        ws = wb.active

        for sheet in self.sheets:
            if sheet == self.sheets[0]:
                ws = wb.active
                ws.title = sheet.name
            else:
                ws = wb.create_sheet(title=sheet.name)
            sheet.dump_to_xls(ws)

        wb.save(xls)




class PulpObjects(object):
    def __init__(self, db_import=False, import_tests=False, db_info=None,
                 db_env=[]):
        self.pulp_sdks = {}
        self.tests = {}
        self.topTests = []
        self.builds = []
        self.runs = []
        self.build_for_id = {}
        self.db_info = db_info
        self.db_env = db_env
        if db_import:
            self.db = plpdb.PulpDb()
            self.builds += self.db.get_builds()
            if import_tests:
                self.runs += self.db.get_tests(self)

            self.build_ids = []
            for build in self.db.get_builds():
                self.build_ids.append(build.id)
                self.build_for_id[build.id] = build
        else:
            self.db = None

    def get_build_from_id(self, id):
        return self.build_for_id[id]

    def get_env_dict(self, env):
        result = {}
        for env_item in env:
            if env_item.find('=') != -1:
                key, value = env_item.split('=')
            else:
                key = env_item
                value = True
            result[key] = value
        return result
                                

    def get_builds(self, branches=[], config=None, env=[]):
        env = self.get_env_dict(env)
        builds = []
        for build in self.db.get_builds():
            if len(branches) == 0 or build.branch in branches:
                if config == None or config in build.config:
                    if len(env.keys()) != 0:
                        build_env = self.get_env_dict(build.env)
                        skip = False
                        for key, value in env.items():
                            if build_env.get(key) != value:
                                skip = True
                                break
                        if skip:
                            continue

                    builds.append(build.id)

        return builds


    def get_configs(self):

        configs = []

        for test in self.topTests:
            for config in test.get_configs():
                if not config in configs:
                    configs.append(config)

        return configs

    def check_build_reg(self, config, ref_build, build):

        for test in self.topTests:
            if not test.check_reg(config, ref_build, build):
                return False

        return True

    def dump_commit(self, header, report, build):
        version = self.get_build_from_id(build).commit_version
        report.add_line(header + version)
        #cmd = ('git show -s --color=never %s' % version).split(' ')
        #result = subprocess.check_output(cmd)
        #report.add_line(result.decode('utf-8'))



    def dump_build_reg_report(self, config, ref_build, build, report):

        report.add_line('')
        report.add_line('')
        report.add_line('Configuration: %s' % config)

        report.add_line('  Build:           %d' % build)
        report.add_line('  Reference build: %d' % ref_build)
        self.dump_commit('  Commit: ', report, build)
        self.dump_commit('  Reference commit: ', report, ref_build)
        report.add_line('')

        table = [['test', str(build)]]

        for test in self.topTests:
            test.dump_reg(config, ref_build, build, table)

        config_name = plptree.get_config_name(config)
        report.add_table('regressions_' + config_name.replace('=','_'), table)

        report.add_line('')


    def check_reg(self, branches, user_config=None, sheet=None, url=None,
                  test_build=None, env=[]):

        if test_build is not None:
            test_build = int(test_build)

        report = None
        if sheet is None:
            report = Report()
            sheet = report.alloc_sheet('Regression report', 'regressions')

        if url is not None:
            sheet.add_line('')
            sheet.add_line('Build URL: ' + url)


        sheet.add_line('')
        sheet.add_line('For each configuration, the last build is compared to the reference build.')
        sheet.add_line('The reference build is the last one which did not add any regression.')
        sheet.add_line('Each test is reporting 4 numbers : regs / fixed / removed / new')
        sheet.add_line('  regs:    Number of regressions, e.g. number of tests which were passing and are now failing')
        sheet.add_line('  fixed:   Number of fixed tests, e.g. which were failing and are now passing')
        sheet.add_line('  removed: Number of tests removed, e.g. which do not exist anymore')
        sheet.add_line('  new:     Number of new tests, e.g. which did not exist')
        sheet.add_line('')

        has_reg = False

        for config in self.get_configs():

            if user_config is not None and config not in user_config:
                continue

            builds = self.get_builds(branches=branches, config=config, env=env)

            # If we don't have at least 2 builds, there is not regression
            if len(builds) < 2: 
                continue

            # Now find out the last build with no regression
            ref_build = builds[0]

            for build in builds[1:]:

                if self.check_build_reg(config, ref_build, build):

                    ref_build = build

            # Now check our build against the ref one
            has_reg = not self.check_build_reg(config, ref_build, builds[-1])

            # And generate the regression report
            if test_build is None:
                test_build = builds[-1]

            self.dump_build_reg_report(config, ref_build, test_build, sheet)

        if report is not None:
            report.dump_to_console()

        return has_reg



    def dump_tests(self, branches, xls=None, user_config=None, mail=None,
                   url=None, author_email=None, build=None, env=[]):

        report = Report()

        sheet = report.alloc_sheet('Regression report', 'regressions')

        self.check_reg(branches, user_config, sheet, url, build, env=env)


        
        sheet = report.alloc_sheet('Test result', 'tests')

        for config in self.get_configs():

            if user_config is not None:
                if config == '' or config not in user_config:
                    continue

            builds = self.get_builds(branches=branches, config=config, env=env)

            columns = ['test']

            for build in builds:
                columns.append(str(build))

            msg = 'Configuration: ' + plptree.get_config_name(config)
            sheet.add_line(msg)

            data = []
            data.append(columns)

            for test in self.topTests:
                test.dump_status(config, builds, table=data)

            config_name = plptree.get_config_name(config)
            sheet.add_table("tests_results_" + config_name.replace('=','_'),
                data)




        report.dump_to_console()

        attachments = []

        if xls is not None:
            report.dump_to_xls(xls)
            attachments.append(xls)

        if mail is not None:
            try:
                report.dump_to_mail(mail, attachments, author_email)
            except:
                print ('Did not manage to send email')

        return


    def drop_all(self):
        self.db.drop_all()
        
    def dump(self, builds=False, tests=False, branches=[]):
        if builds:
            x = PrettyTable([
                'id', 'package', 'config', 'version', 'branch', 'status', 'start_date',
                'end_date', 'artifact', 'deps', 'tests', 'passed', 'testrun'
            ])
            x.align = "l"
            x.align['id'] = "c"
            for build in self.builds:
                if len(branches) == 0 or \
                      build.branch in branches:
                    x.add_row([
                        build.id,
                        build.package, build.config, build.version, build.branch, build.status,
                        build.start_date, build.end_date, build.artifact,
                        build.deps, build.tests, build.passed, build.testrun])

            print(x)

        if tests:

            x = PrettyTable([
                'name', 'config', 'status', 'time', 'build', 'metrics'
            ])
            x.align = "l"
            x.align['build'] = "c"
            x.align['time'] = "r"
            x.float_format = '.2'
            for test in self.runs:

                if len(branches) != 0:
                    build = self.build_for_id[test.build]
                    if not build.branch in branches:
                        continue


                x.add_row([
                    test.test.name, test.config, test.status, test.duration,
                    test.build, test.metrics
                ])

            print(x)

    def getNbTests(self):
        result = 0
        for test in self.topTests:
            result += test.getNbTests()
        return result

    def getNbSuccess(self):
        result = 0
        for test in self.topTests:
            result += test.getNbSuccess()
        return result

    def getTest(self, name):
        if self.tests.get(name) is None:
            nameList = name.split(':')
            parent = None
            if len(nameList) > 1:
                names = ':'.join(nameList[0:len(nameList) - 1])
                parent = self.getTest(names)
            self.tests[name] = Test(name, parent)
            if len(nameList) == 1:
                self.topTests.append(self.tests[name])
        return self.tests.get(name)

    def dumpToFile(self, path):
        pass

    def loadFromFile(self, path):
        pass

    def mergeTo(self, objects):
        pass

    def dumpTestsToConsole(self):
        x = PrettyTable(['test', 'config', 'time', 'passed/total'])
        x.align["test"] = "l"
        x.align["config"] = "l"
        for test in self.topTests:
            test.dumpToConsole(x)
        print()
        print(x)

    def dumpTestsToJunit(self, reportPath):
        try:
            os.makedirs(reportPath)
        except:
            pass

        for test in self.topTests:
            test.dumpToJunit(reportPath)


def PulpSdk_get(version):
    return pulp_sdks.get(version)
